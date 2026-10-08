"""APA 7 manuscripts split a caption into a label line ("Figure 1") and a
title line. TME wants one centered line ("Figure 1. Title") and centered
figure images."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import fixup
from tme_template.styles import (
    register_body_style, register_title_style,
    register_heading_styles, register_remaining_styles,
)


def _make_doc_with_styles():
    doc = Document()
    register_body_style(doc)
    register_title_style(doc)
    register_heading_styles(doc)
    register_remaining_styles(doc)
    return doc


def _add_drawing_paragraph(doc, style="TME Body", anchored=False):
    """Image-only paragraph. Inline by default (wp:inline); anchored=True makes
    a floating picture (wp:anchor), whose paragraph alignment means nothing."""
    p = doc.add_paragraph("", style=style)
    drawing = OxmlElement("w:drawing")
    drawing.append(OxmlElement("wp:anchor" if anchored else "wp:inline"))
    p.add_run()._r.append(drawing)
    return p


def _add_table_after(doc):
    return doc.add_table(rows=1, cols=1)


def _texts(doc):
    return [p.text for p in doc.paragraphs]


# ---- merge_split_captions ----

def test_label_and_title_lines_merge_into_one_figure_caption():
    doc = _make_doc_with_styles()
    doc.add_paragraph("Body before.", style="TME Body")
    label = doc.add_paragraph("Figure 1", style="TME Body")  # source style varies; text is the signal
    doc.add_paragraph("UTG for the Cake Task, Adapted from Norton et al. (2023)", style="TME Figure Caption")
    _add_drawing_paragraph(doc)

    stats = fixup.merge_split_captions(doc)

    assert stats == {"labels_restyled": 1, "titles_merged": 1}
    assert _texts(doc) == [
        "Body before.",
        "Figure 1. UTG for the Cake Task, Adapted from Norton et al. (2023)",
        "",
    ]
    assert label.style.name == "TME Figure Caption"


def test_table_label_with_trailing_period_merges_into_table_caption():
    doc = _make_doc_with_styles()
    label = doc.add_paragraph("Table 2.", style="TME Body")
    doc.add_paragraph("Scoring Rubric", style="TME Body")

    fixup.merge_split_captions(doc)

    assert _texts(doc) == ["Table 2. Scoring Rubric"]
    assert label.style.name == "TME Table Caption"


def test_colon_after_label_is_normalised_to_period():
    doc = _make_doc_with_styles()
    doc.add_paragraph("Figure 3:", style="TME Body")
    doc.add_paragraph("Sidewalk Task", style="TME Body")
    fixup.merge_split_captions(doc)
    assert _texts(doc) == ["Figure 3. Sidewalk Task"]


def test_title_run_formatting_survives_the_merge():
    doc = _make_doc_with_styles()
    doc.add_paragraph("Figure 1", style="TME Body")
    title = doc.add_paragraph(style="TME Body")
    title.add_run("Graph of ")
    title.add_run("f").italic = True
    _add_drawing_paragraph(doc)

    fixup.merge_split_captions(doc)

    merged = doc.paragraphs[0]
    assert merged.text == "Figure 1. Graph of f"
    assert [r.text for r in merged.runs if r.italic] == ["f"]


def test_label_followed_directly_by_image_is_restyled_but_not_merged():
    doc = _make_doc_with_styles()
    label = doc.add_paragraph("Figure 4", style="TME Body")
    _add_drawing_paragraph(doc)

    stats = fixup.merge_split_captions(doc)

    assert stats == {"labels_restyled": 1, "titles_merged": 0}
    assert label.style.name == "TME Figure Caption"
    assert _texts(doc) == ["Figure 4", ""]


def test_label_followed_by_heading_is_not_merged():
    doc = _make_doc_with_styles()
    doc.add_paragraph("Figure 5", style="TME Body")
    doc.add_paragraph("Results", style="TME H1")
    fixup.merge_split_captions(doc)
    assert _texts(doc) == ["Figure 5", "Results"]


def test_body_sentence_starting_with_figure_is_left_alone():
    doc = _make_doc_with_styles()
    body = doc.add_paragraph("Figure 1 depicts a UTG of a hypothetical solution.", style="TME Body")
    doc.add_paragraph("Another body paragraph.", style="TME Body")
    stats = fixup.merge_split_captions(doc)
    assert stats == {"labels_restyled": 0, "titles_merged": 0}
    assert body.style.name == "TME Body"


def test_one_line_caption_is_left_alone():
    doc = _make_doc_with_styles()
    doc.add_paragraph("Figure 1. Already one line.", style="TME Figure Caption")
    _add_drawing_paragraph(doc)
    stats = fixup.merge_split_captions(doc)
    assert stats == {"labels_restyled": 0, "titles_merged": 0}


def test_empty_spacer_between_label_and_title_is_removed():
    doc = _make_doc_with_styles()
    doc.add_paragraph("Figure 6", style="TME Body")
    doc.add_paragraph("", style="TME Body")
    doc.add_paragraph("Packages Task", style="TME Body")
    _add_drawing_paragraph(doc)
    fixup.merge_split_captions(doc)
    assert _texts(doc) == ["Figure 6. Packages Task", ""]


def test_table_label_followed_by_a_table_does_not_eat_the_note_below_it():
    doc = _make_doc_with_styles()
    label = doc.add_paragraph("Table 1", style="TME Body")
    _add_table_after(doc)
    note = doc.add_paragraph("Note. Values are means.", style="TME Body")

    stats = fixup.merge_split_captions(doc)

    assert stats == {"labels_restyled": 1, "titles_merged": 0}
    assert label.style.name == "TME Table Caption"
    assert _texts(doc) == ["Table 1", "Note. Values are means."]
    assert note.style.name == "TME Body"


def test_title_set_entirely_in_run_italics_comes_out_plain():
    """APA 7 titles are italic as a whole; that is style, not emphasis."""
    doc = _make_doc_with_styles()
    doc.add_paragraph("Figure 1", style="TME Body")
    title = doc.add_paragraph(style="TME Body")
    title.add_run("UTG for the ").italic = True
    title.add_run("Cake Task").italic = True
    _add_drawing_paragraph(doc)

    fixup.merge_split_captions(doc)

    merged = doc.paragraphs[0]
    assert merged.text == "Figure 1. UTG for the Cake Task"
    assert not any(r.italic for r in merged.runs)


def test_trailing_blank_run_on_label_does_not_leave_a_space_before_the_period():
    doc = _make_doc_with_styles()
    label = doc.add_paragraph(style="TME Body")
    label.add_run("Figure ")
    label.add_run("1")
    label.add_run(" ")
    doc.add_paragraph("Cake Task", style="TME Body")

    fixup.merge_split_captions(doc)

    assert _texts(doc) == ["Figure 1. Cake Task"]


def test_page_break_paragraph_between_label_and_title_is_kept():
    doc = _make_doc_with_styles()
    doc.add_paragraph("Figure 2", style="TME Body")
    brk = doc.add_paragraph(style="TME Body")
    brk.add_run().add_break(__import__("docx.enum.text", fromlist=["WD_BREAK"]).WD_BREAK.PAGE)
    doc.add_paragraph("Title after break", style="TME Body")

    stats = fixup.merge_split_captions(doc)

    assert stats == {"labels_restyled": 1, "titles_merged": 0}
    assert len(doc.paragraphs) == 3
    assert brk._p.find(".//" + qn("w:br")) is not None


def test_section_break_paragraph_between_label_and_title_is_kept():
    doc = _make_doc_with_styles()
    doc.add_paragraph("Figure 2", style="TME Body")
    sect = doc.add_paragraph(style="TME Body")
    sect._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    doc.add_paragraph("Title after break", style="TME Body")

    fixup.merge_split_captions(doc)

    assert len(doc.paragraphs) == 3
    assert sect._p.pPr.find(qn("w:sectPr")) is not None


def test_unnumbered_label_is_not_folded_into_a_previous_label():
    """Two SEQ labels with uncached results read as bare 'Figure' lines."""
    doc = _make_doc_with_styles()
    doc.add_paragraph("Figure", style="TME Body")
    doc.add_paragraph("Figure", style="TME Body")
    _add_drawing_paragraph(doc)

    stats = fixup.merge_split_captions(doc)

    assert stats["titles_merged"] == 0
    assert _texts(doc) == ["Figure", "Figure", ""]


# ---- center_image_paragraphs ----

def test_anchored_floating_image_is_not_centered_or_counted():
    doc = _make_doc_with_styles()
    img = _add_drawing_paragraph(doc, anchored=True)
    assert fixup.center_image_paragraphs(doc) == 0
    assert img.alignment is None


def test_images_before_the_body_section_break_are_left_alone():
    doc = _make_doc_with_styles()
    cover_img = _add_drawing_paragraph(doc)
    brk = doc.add_paragraph(style="TME Body")
    brk._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    body_img = _add_drawing_paragraph(doc)

    assert fixup.center_image_paragraphs(doc) == 1
    assert cover_img.alignment is None
    assert body_img.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_image_only_paragraph_is_centered_whatever_its_style():
    doc = _make_doc_with_styles()
    doc.styles.add_style("PictureANDLetterLabels", 1)  # WD_STYLE_TYPE.PARAGRAPH
    img = _add_drawing_paragraph(doc, style="PictureANDLetterLabels")

    n = fixup.center_image_paragraphs(doc)

    assert n == 1
    assert img.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_text_paragraph_with_inline_image_is_not_centered():
    doc = _make_doc_with_styles()
    p = doc.add_paragraph("An inline icon ", style="TME Body")
    p.add_run()._r.append(OxmlElement("w:drawing"))
    n = fixup.center_image_paragraphs(doc)
    assert n == 0
    assert p.alignment is None


def test_already_centered_image_is_not_counted():
    doc = _make_doc_with_styles()
    img = _add_drawing_paragraph(doc)
    img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    assert fixup.center_image_paragraphs(doc) == 0


# ---- run_fixup wiring ----

def test_run_fixup_merges_and_centers(tmp_path):
    doc = _make_doc_with_styles()
    doc.add_paragraph("Figure 1", style="TME Body")
    doc.add_paragraph("Cake Task", style="TME Body")
    _add_drawing_paragraph(doc)
    path = tmp_path / "t.docx"
    doc.save(path)

    stats = fixup.run_fixup(str(path))

    out = Document(path)
    assert [p.text for p in out.paragraphs] == ["Figure 1. Cake Task", ""]
    assert out.paragraphs[0].style.name == "TME Figure Caption"
    assert out.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert stats["split_captions"] == {"labels_restyled": 1, "titles_merged": 1}
    assert stats["images_centered"] == 1
