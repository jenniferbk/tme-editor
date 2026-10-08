"""Metadata parsing is lenient about nulls and the whole manuscript is read."""
from datetime import date

from docx import Document

import extractor
from extractor import ArticleMeta, extract_manuscript_text, to_article_meta


def test_nulls_and_bad_indices_become_safe_defaults():
    meta = to_article_meta({
        "title": None,
        "authors": [{"name": None, "affiliation_num": 7, "bio": None, "email": None, "corresponding": None},
                    "not a dict"],
        "affiliations": ["X"],
        "abstract": None, "keywords": None,
        "received": None, "revised": None, "accepted": None, "published": None, "doi": None,
    })
    assert meta.title == "" and meta.abstract == "" and meta.keywords == []
    assert len(meta.authors) == 1
    assert meta.authors[0].name == "" and meta.authors[0].affiliation_num == 1
    assert meta.authors[0].corresponding is False


def test_year_defaults_to_this_year():
    assert ArticleMeta().year == date.today().year


def test_whole_manuscript_is_read(tmp_path):
    path = tmp_path / "long.docx"
    doc = Document()
    for i in range(400):
        doc.add_paragraph("x" * 100 + f" {i}")
    doc.add_paragraph("AUTHOR BIO AT THE END")
    doc.save(path)
    assert extract_manuscript_text(str(path)).endswith("AUTHOR BIO AT THE END")


def test_response_schema_is_passed_to_gemini(monkeypatch):
    captured = {}

    class _Resp:
        text = '{"title": "T", "authors": [], "affiliations": [], "abstract": "", "keywords": [], "received": "", "revised": "", "accepted": "", "published": "", "doi": ""}'

    class _Models:
        def generate_content(self, **kw):
            captured.update(kw)
            return _Resp()

    class _Client:
        def __init__(self, api_key):
            self.models = _Models()

    monkeypatch.setattr(extractor.genai, "Client", _Client)
    meta = extractor.extract_metadata("text", api_key="k")
    assert meta.title == "T"
    assert captured["config"]["response_schema"] is extractor.MetadataSchema
