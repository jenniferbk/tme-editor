"""Fixup heuristics: block quotes need length, caption swap is index-exact,
header rows are only marked when they look like headers."""
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import fixup
from tme_template.styles import (
    register_body_style, register_title_style,
    register_heading_styles, register_remaining_styles,
)


def _doc():
    doc = Document()
    for f in (register_body_style, register_title_style, register_heading_styles, register_remaining_styles):
        f(doc)
    return doc


def _image_paragraph(doc):
    p = doc.add_paragraph("", style="TME Body")
    d = OxmlElement("w:drawing")
    d.append(OxmlElement("wp:inline"))
    p.add_run()._r.append(d)
    return p


def test_short_indented_task_part_is_not_a_block_quote():
    doc = _doc()
    p = doc.add_paragraph("(a) Sketch the graph of f on [0, 4].", style="TME Body")
    p.paragraph_format.left_indent = Pt(36)
    assert fixup.remap_block_quotes(doc) == 0
    assert p.style.name == "TME Body"


def test_long_indented_quotation_is_a_block_quote():
    doc = _doc()
    p = doc.add_paragraph(" ".join(["word"] * 45), style="TME Body")
    p.paragraph_format.left_indent = Pt(36)
    assert fixup.remap_block_quotes(doc) == 1
    assert p.style.name == "TME Block Quote"


def test_swap_moves_two_identically_prefixed_captions_correctly():
    doc = _doc()
    for n in (1, 2):
        _image_paragraph(doc)
        doc.add_paragraph(f"Figure {n}. Same prefix caption", style="TME Figure Caption")
    report = fixup.report_below_element_captions(doc)
    assert [r["index"] for r in report] == [1, 3]

    moved = fixup.swap_captions_above(doc, report)

    assert moved == 2
    texts = [p.text for p in doc.paragraphs]
    assert texts == ["Figure 1. Same prefix caption", "", "Figure 2. Same prefix caption", ""]


def test_swap_ignores_a_figure_entry_whose_predecessor_is_text():
    doc = _doc()
    doc.add_paragraph("Just text.", style="TME Body")
    doc.add_paragraph("Figure 1. Not below an image", style="TME Figure Caption")
    fake = [{"index": 1, "kind": "figure", "preview": "Figure 1."}]
    assert fixup.swap_captions_above(doc, fake) == 0


def test_header_row_only_when_it_looks_like_one():
    doc = _doc()
    one_row = doc.add_table(rows=1, cols=2)
    one_row.cell(0, 0).text = "a"
    layout = doc.add_table(rows=2, cols=2)          # first row has an empty cell
    layout.cell(0, 0).text = "img"
    data = doc.add_table(rows=2, cols=2)
    for c in data.rows[0].cells:
        c.text = "head"

    fixup.fix_content_tables(doc, skip_indices=())

    def has_header(t):
        return t._tbl.findall(qn("w:tr"))[0].find(qn("w:trPr") + "/" + qn("w:tblHeader")) is not None
    assert not has_header(one_row) and not has_header(layout) and has_header(data)
