"""The top-of-body cleaner removes pasted cover duplicates and nothing else."""
from docx import Document

import apply_styles
import classifier
from extractor import ArticleMeta, AuthorMeta
from tme_template.styles import (
    register_body_style, register_title_style,
    register_heading_styles, register_remaining_styles,
)

META = ArticleMeta(
    title="Integration by Substitution: An Emergent Quantitative Reasoning Approach to U-Substitution",
    authors=[AuthorMeta(name="Kevin C. Moore")],
    affiliations=["Department of Mathematics and Science Education, University of Georgia"],
    abstract="We present a conceptual analysis of integration by substitution grounded in quantitative reasoning.",
)


def _dup(text):
    return apply_styles._looks_like_cover_duplicate(text, META)


def test_real_body_openers_survive():
    assert not _dup("Quantitative Reasoning")                                   # heading inside the title
    assert not _dup("Published research on u-substitution has largely focused on procedures.")
    assert not _dup("Department of Mathematics faculty participated in the study.")
    assert not _dup("Accepted practice in calculus instruction emphasizes procedures.")
    assert not _dup("Published in 2019, Smith found X.")
    assert not _dup("Published research from 2019 shows a gap.")
    assert not _dup("Revised version 2 of the task was harder.")


def test_genuine_cover_duplicates_are_caught():
    assert _dup(META.title)
    assert _dup(META.title + "1")                                              # with affiliation mark
    assert _dup("Abstract")
    assert _dup("Keywords: calculus, quantitative reasoning")
    assert _dup("Received: March 3, 2026")
    assert _dup("Accepted 12 May 2026")
    assert _dup("Published online June 2026")
    assert _dup("Received: 2026-03-03")
    assert _dup("Revised 3/3/2026")
    assert _dup(META.abstract)
    assert _dup("Kevin C. Moore")
    assert _dup(META.affiliations[0])
    assert _dup("† Corresponding author: kvcmoore@uga.edu")


def test_plain_document_without_title_classifies_from_paragraph_zero(tmp_path, monkeypatch):
    doc = Document()
    for f in (register_body_style, register_title_style,
              register_heading_styles, register_remaining_styles):
        f(doc)
    doc.add_paragraph("First paragraph of a plain document.")
    doc.add_paragraph("Second paragraph of a plain document.")
    path = tmp_path / "plain.docx"
    doc.save(path)

    monkeypatch.setattr(
        classifier, "classify_paragraphs_with_report",
        lambda texts, **kw: (["body"] * len(texts), {"missing": 0, "invalid": 0, "duplicates": 0}),
    )
    apply_styles.apply_styles(str(path), ArticleMeta(title="Zzz"))

    styles = [p.style.name for p in Document(str(path)).paragraphs if p.text.strip()]
    assert styles == ["TME Body", "TME Body"]
