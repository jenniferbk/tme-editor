"""Build a starter, paste a synthetic body, run both phases with a fixed
classifier, and check every paragraph's style and the untouched cover."""
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import classifier
from apply_styles import apply_styles
from article_starter import build_article_starter
from extractor import ArticleMeta, AuthorMeta
from fixup import run_fixup

HYPERLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"

META = ArticleMeta(
    title="Expanding the Scope of Unit Transformation Graphs",
    authors=[AuthorMeta(name="Joseph Antonides", bio="bio"), AuthorMeta(name="Anderson Norton", bio="bio")],
    affiliations=["Virginia Tech"], abstract="We extend UTGs.", keywords=["UTG"], year=2026,
)

LABELS = {
    "Introduction": "heading_1",
    "Figure 1 depicts a UTG of a hypothetical solution to the Cake Task.": "body",
    "Table 2 presents the coding scheme used in this study.": "body",
    "Figure 1": "skip",
    "UTG for the Cake Task": "caption_figure",
    "Students' Mathematics": "heading_2",
    "Antonides, J., & Norton, A. (2023). Units. Journal, 1(1), 1–2. https://doi.org/10.1/x": "reference",
    "(a) Sketch the graph.": "body",
}


def _fake_classifier(texts, **kwargs):
    return [LABELS.get(t.strip(), "body") for t in texts], {"missing": 0, "invalid": 0, "duplicates": 0}


def _hyperlink(p, text):
    r_id = p.part.relate_to("https://doi.org/10.1/x", HYPERLINK_REL, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr"); f = OxmlElement("w:rFonts")
    f.set(qn("w:ascii"), "Times New Roman"); rPr.append(f); r.append(rPr)
    t = OxmlElement("w:t"); t.text = text; r.append(t); h.append(r); p._p.append(h)
    return r


def _paste_body(path):
    doc = Document(path)
    doc.add_paragraph("Expanding the Scope of Unit Transformation Graphs")      # pasted title (duplicate)
    doc.add_paragraph("Introduction")
    doc.add_paragraph("Figure 1 depicts a UTG of a hypothetical solution to the Cake Task.")
    doc.add_paragraph("Table 2 presents the coding scheme used in this study.")
    doc.add_paragraph("Figure 1")
    doc.add_paragraph("UTG for the Cake Task")
    img = doc.add_paragraph()
    d = OxmlElement("w:drawing"); d.append(OxmlElement("wp:inline")); img.add_run()._r.append(d)
    doc.add_paragraph("Students' Mathematics")
    ref = doc.add_paragraph("Antonides, J., & Norton, A. (2023). Units. Journal, 1(1), 1–2. ")
    ref.paragraph_format.first_line_indent = Pt(-18); ref.paragraph_format.left_indent = Pt(18)
    _hyperlink(ref, "https://doi.org/10.1/x")
    task = doc.add_paragraph("(a) Sketch the graph.")
    task.paragraph_format.left_indent = Pt(36)
    t = doc.add_table(rows=2, cols=2)
    for c in t.rows[0].cells: c.text = "Head"
    t.cell(1, 0).text = "v"
    doc.save(path)


def test_golden_pipeline(tmp_path, monkeypatch):
    monkeypatch.setattr(classifier, "classify_paragraphs_with_report", _fake_classifier)
    path = tmp_path / "proof.docx"
    build_article_starter(meta=META, headshots={}, out_path=path)
    card_xml_before = Document(path).tables[2]._tbl.xml
    _paste_body(path)

    style_stats = apply_styles(str(path), META)
    fixup_stats = run_fixup(str(path))

    doc = Document(path)
    body = {p.text: p for p in doc.paragraphs}
    assert style_stats["classifier"] == "gemini"
    assert style_stats["deleted_preamble_previews"] == ["Expanding the Scope of Unit Transformation Graphs"]
    assert body["Introduction"].style.name == "TME H1"
    assert body["Students' Mathematics"].style.name == "TME H2"
    assert body["Figure 1 depicts a UTG of a hypothetical solution to the Cake Task."].style.name == "TME Body"
    assert body["Table 2 presents the coding scheme used in this study."].style.name == "TME Body"
    assert body["Figure 1. UTG for the Cake Task"].style.name == "TME Figure Caption"
    assert body["(a) Sketch the graph."].style.name == "TME Body"
    img_para = [p for p in doc.paragraphs if p._p.find(".//" + qn("w:drawing")) is not None][-1]
    assert img_para.alignment == WD_ALIGN_PARAGRAPH.CENTER
    ref_para = next(p for p in doc.paragraphs if p.text.startswith("Antonides, J."))
    assert ref_para.style.name == "TME Reference"
    link_runs = ref_para._p.findall(".//" + qn("w:hyperlink") + "/" + qn("w:r"))
    assert link_runs and all(r.find(qn("w:rPr") + "/" + qn("w:rFonts")) is None for r in link_runs)
    assert doc.tables[2]._tbl.xml == card_xml_before            # author card untouched
    content = doc.tables[3]
    assert all(p.style.name == "TME Table Text" for row in content.rows for c in row.cells for p in c.paragraphs)
    assert content._tbl.findall(qn("w:tr"))[0].find(qn("w:trPr") + "/" + qn("w:tblHeader")) is not None
    assert fixup_stats["split_captions"] == {"labels_restyled": 1, "titles_merged": 1}
