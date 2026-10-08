"""Only a real caption label ("Figure 1." / "Table 2:") makes a paragraph a
caption. Body sentences that merely start with the word are left alone."""
from docx import Document

import apply_styles
import fixup
from tme_template.styles import (
    register_body_style, register_title_style,
    register_heading_styles, register_remaining_styles,
)

SENTENCES = [
    "Figure 1 depicts a UTG of a hypothetical solution.",
    "Table 2 presents the coding scheme used in this study.",
    "Figures 3 and 4 show the two students' graphs.",
    "Tables were cleared before the second interview.",
]
CAPTIONS = ["Figure 1. A real caption.", "Table 2: Participants", "Figure A1. Appendix figure"]


def _doc():
    doc = Document()
    for f in (register_body_style, register_title_style, register_heading_styles, register_remaining_styles):
        f(doc)
    return doc


def test_fixup_leaves_body_sentences_alone_but_catches_captions(tmp_path):
    doc = _doc()
    for s in SENTENCES + CAPTIONS:
        doc.add_paragraph(s, style="TME Body")
    path = tmp_path / "t.docx"
    doc.save(path)

    fixup.run_fixup(str(path))

    out = {p.text: p.style.name for p in Document(path).paragraphs}
    assert all(out[s] == "TME Body" for s in SENTENCES), out
    assert out["Figure 1. A real caption."] == "TME Figure Caption"
    assert out["Table 2: Participants"] == "TME Table Caption"
    assert out["Figure A1. Appendix figure"] == "TME Figure Caption"


def test_heuristic_classifier_uses_the_same_rule():
    doc = _doc()
    p = doc.add_paragraph("x", style="TME Body")
    for s in SENTENCES:
        assert apply_styles._heuristic_classify(s, p) == "body", s
    assert apply_styles._heuristic_classify("Figure 1. A real caption.", p) == "caption_figure"
    assert apply_styles._heuristic_classify("Table 2: Participants", p) == "caption_table"


def test_swapped_caption_styles_are_still_corrected():
    doc = _doc()
    fig = doc.add_paragraph("Figure 1. Wrong style.", style="TME Table Caption")
    tab = doc.add_paragraph("Table 1. Wrong style.", style="TME Figure Caption")
    fixup.fix_caption_classifications(doc)
    assert fig.style.name == "TME Figure Caption"
    assert tab.style.name == "TME Table Caption"
