"""ArticleMeta rides inside the starter so Phase 2 can recover it."""
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from extractor import ArticleMeta, AuthorMeta
from session_meta import DOCVAR_NAME, choose_meta, embed_meta, meta_from_json, meta_to_json, read_meta


def _meta():
    return ArticleMeta(title="T", authors=[AuthorMeta(name="Ada Lovelace", corresponding=True)],
                       affiliations=["UGA"], abstract="a", keywords=["k"], volume=34, number=2)


def test_round_trip_through_json():
    m = meta_from_json(meta_to_json(_meta()))
    assert m == _meta()


def test_unknown_keys_from_a_newer_starter_are_ignored():
    s = meta_to_json(_meta()).replace('"title": "T"', '"title": "T", "future_field": 1')
    assert meta_from_json(s).title == "T"


def test_embed_and_read_back_long_metadata_survives_resave(tmp_path):
    m = _meta()
    m.abstract = "Abstract with \"quotes\" & <tags> é. " * 100   # about 3,500 chars
    m.authors[0].bio = "Bio – with dashes. " * 80                # about 1,500 chars
    assert len(m.abstract) >= 3000 and len(m.authors[0].bio) >= 1500
    path = tmp_path / "s.docx"
    Document().save(path)
    embed_meta(path, m)
    assert read_meta(path) == m
    Document(str(path)).save(str(path))      # a Word-style resave must keep the variable
    assert read_meta(path) == m


def test_embedding_twice_replaces_the_value(tmp_path):
    path = tmp_path / "s.docx"
    Document().save(path)
    embed_meta(path, _meta())
    m2 = _meta()
    m2.title = "Second"
    embed_meta(path, m2)
    assert read_meta(path).title == "Second"
    root = Document(str(path)).settings.element
    assert len(root.findall(qn("w:docVars") + "/" + qn("w:docVar"))) == 1


def test_corrupt_marker_reads_none(tmp_path):
    path = tmp_path / "bad.docx"
    doc = Document()
    docvars = OxmlElement("w:docVars")
    var = OxmlElement("w:docVar")
    var.set(qn("w:name"), DOCVAR_NAME)
    var.set(qn("w:val"), "not json")
    docvars.append(var)
    doc.settings.element.append(docvars)
    doc.save(path)
    assert read_meta(path) is None


def test_document_without_meta_reads_none(tmp_path):
    path = tmp_path / "plain.docx"
    Document().save(path)
    assert read_meta(path) is None


def test_oversized_metadata_raises(tmp_path):
    m = _meta()
    m.abstract = "x" * 70_000
    path = tmp_path / "big.docx"
    Document().save(path)
    with pytest.raises(ValueError, match="too large"):
        embed_meta(path, m)


def test_choose_meta_embedded_only():
    chosen, differs = choose_meta(_meta(), None)
    assert chosen == _meta() and differs is False


def test_choose_meta_session_only():
    chosen, differs = choose_meta(None, _meta())
    assert chosen == _meta() and differs is False
    assert choose_meta(None, None) == (None, False)


def test_choose_meta_both_differing_prefers_embedded():
    other = _meta()
    other.title = "Article A"
    chosen, differs = choose_meta(_meta(), other)
    assert chosen.title == "T" and differs is True
    assert choose_meta(_meta(), _meta()) == (_meta(), False)
