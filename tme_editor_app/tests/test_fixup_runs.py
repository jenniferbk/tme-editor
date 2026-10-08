"""Run-level stripping reaches hyperlinks; tables get a TME style; footnote
font rewrite leaves symbol fonts alone; the redundant reference strip is gone."""
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import fixup
from tme_template.styles import (
    register_body_style, register_title_style,
    register_heading_styles, register_remaining_styles,
)

HYPERLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


def _doc():
    doc = Document()
    for f in (register_body_style, register_title_style, register_heading_styles, register_remaining_styles):
        f(doc)
    return doc


def _add_hyperlink_run(p, text, font="Times New Roman"):
    r_id = p.part.relate_to("https://doi.org/10.1/x", HYPERLINK_REL, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), font)
    rPr.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)
    return r


def test_strip_direct_formatting_reaches_runs_inside_hyperlinks():
    doc = _doc()
    p = doc.add_paragraph("Smith, J. (2020). Title. ", style="TME Reference")
    r = _add_hyperlink_run(p, "https://doi.org/10.1/x")

    fixup.strip_direct_formatting(doc)

    rPr = r.find(qn("w:rPr"))
    assert rPr.find(qn("w:rFonts")) is None and rPr.find(qn("w:sz")) is None


def test_reference_with_only_a_hyperlink_run_is_still_stripped():
    doc = _doc()
    p = doc.add_paragraph(style="TME Reference")
    r = _add_hyperlink_run(p, "https://doi.org/10.1/x")
    fixup.strip_direct_formatting(doc)
    assert r.find(qn("w:rPr")).find(qn("w:rFonts")) is None


def test_redundant_reference_strip_is_gone():
    assert not hasattr(fixup, "strip_reference_run_formatting")


def test_content_table_cells_get_table_text_style():
    doc = _doc()
    t = doc.add_table(rows=1, cols=1)
    cell_p = t.cell(0, 0).paragraphs[0]
    cell_p.add_run("x").font.name = "Times New Roman"

    fixup.update_styles(doc)              # creates TME Table Text if missing
    fixup.normalize_table_cells(doc, skip_indices=())

    assert cell_p.style.name == "TME Table Text"
    assert doc.styles["TME Table Text"].font.name == "Georgia"


def test_footnote_rewrite_keeps_symbol_fonts():
    xml = ('<w:footnotes xmlns:w="x"><w:r><w:rPr><w:rFonts w:ascii="Times New Roman"/>'
           '<w:sz w:val="20"/></w:rPr></w:r><w:r><w:rPr><w:rFonts w:ascii="Symbol" w:hAnsi="Symbol"/>'
           '</w:rPr></w:r><w:r><w:rPr><w:rFonts w:ascii="Cambria Math"/></w:rPr></w:r></w:footnotes>')
    out, stats = fixup.rewrite_footnote_xml(xml)
    assert 'w:ascii="Georgia"' in out
    assert 'w:ascii="Symbol"' in out and 'w:ascii="Cambria Math"' in out
    assert stats["rfonts_rewritten"] == 1 and stats["sz_stripped"] == 1


def test_cell_paragraph_direct_spacing_and_indent_stripped_but_alignment_kept():
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    doc = _doc()
    t = doc.add_table(rows=1, cols=1)
    cell_p = t.cell(0, 0).paragraphs[0]
    cell_p.add_run("1.5")
    cell_p.paragraph_format.space_after = Pt(18)
    cell_p.paragraph_format.left_indent = Pt(20)
    cell_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    fixup.update_styles(doc)
    fixup.normalize_table_cells(doc, skip_indices=())

    pPr = cell_p._p.pPr
    assert pPr.find(qn("w:spacing")) is None
    assert pPr.find(qn("w:ind")) is None
    assert pPr.find(qn("w:jc")).get(qn("w:val")) == "right"


def _add_textbox_run(p, text="label"):
    """A run holding a drawing whose text box contains its own paragraph and
    run, formatted 9pt Arial bold, like a label in a Word-drawn figure."""
    from lxml import etree
    wps = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
    outer = OxmlElement("w:r")
    drawing = OxmlElement("w:drawing")
    txbx = etree.SubElement(drawing, "{%s}txbx" % wps)
    content = etree.SubElement(txbx, qn("w:txbxContent"))
    inner_p = etree.SubElement(content, qn("w:p"))
    inner_r = etree.SubElement(inner_p, qn("w:r"))
    rPr = etree.SubElement(inner_r, qn("w:rPr"))
    fonts = etree.SubElement(rPr, qn("w:rFonts"))
    fonts.set(qn("w:ascii"), "Arial")
    etree.SubElement(rPr, qn("w:b"))
    etree.SubElement(rPr, qn("w:sz")).set(qn("w:val"), "18")
    t = etree.SubElement(inner_r, qn("w:t"))
    t.text = text
    outer.append(drawing)
    p._p.append(outer)
    return inner_r


def test_text_box_labels_are_left_alone_but_direct_runs_are_stripped():
    doc = _doc()
    p = doc.add_paragraph(style="TME Body")
    direct = p.add_run("body text")
    direct.font.name = "Times New Roman"
    inner = _add_textbox_run(p)

    assert len(fixup._iter_runs(p)) == 2          # direct run + the run holding the drawing; the label is excluded

    fixup.strip_direct_formatting(doc)

    inner_rPr = inner.find(qn("w:rPr"))
    for tag in ("w:rFonts", "w:b", "w:sz"):
        assert inner_rPr.find(qn(tag)) is not None, tag
    assert direct._r.find(qn("w:rPr")).find(qn("w:rFonts")) is None


def test_trPr_flags_follow_schema_order_when_trHeight_exists():
    doc = _doc()
    t = doc.add_table(rows=2, cols=1)
    for row in t.rows:
        row.cells[0].text = "x"
    tr = t.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:trHeight"))

    fixup.fix_content_tables(doc, skip_indices=())

    names = [c.tag.split("}")[1] for c in tr.find(qn("w:trPr"))]
    assert names == ["cantSplit", "trHeight", "tblHeader"]
