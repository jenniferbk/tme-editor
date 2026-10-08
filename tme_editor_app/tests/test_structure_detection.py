"""Body start and cover tables are found by the starter's structure, not by
'last section break' or hard-coded table indices."""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

import apply_styles
import fixup
from article_starter import build_article_starter
from extractor import ArticleMeta, AuthorMeta


def _starter(tmp_path):
    out = tmp_path / "starter.docx"
    meta = ArticleMeta(title="A Title", authors=[AuthorMeta(name="Ada Lovelace", bio="b")],
                       affiliations=["UGA"], abstract="abs", keywords=["k"])
    build_article_starter(meta=meta, headshots={}, out_path=out)
    return out


def test_body_starts_at_the_placeholder_even_after_a_pasted_section_break(tmp_path):
    doc = Document(_starter(tmp_path))
    paras = list(doc.paragraphs)
    start = apply_styles.find_body_start_index(paras)
    assert paras[start].text.startswith("[Paste")

    doc.add_paragraph("Body text.")
    doc.add_section(WD_SECTION.NEW_PAGE)      # a landscape section pasted by the editor
    doc.add_paragraph("Wide table lives here.")
    paras = list(doc.paragraphs)
    assert apply_styles.find_body_start_index(paras) == start


def test_plain_document_without_breaks_classifies_everything():
    doc = Document()
    doc.add_paragraph("a")
    assert apply_styles.find_body_start_index(list(doc.paragraphs)) is None


def test_cover_tables_are_counted_by_position(tmp_path):
    doc = Document(_starter(tmp_path))
    assert fixup.cover_table_count(doc) == 3      # masthead, tagline, author card
    doc.add_table(rows=2, cols=2)                 # pasted content table
    assert fixup.cover_table_count(doc) == 3


def test_fix_content_tables_leaves_the_author_card_alone(tmp_path):
    doc = Document(_starter(tmp_path))
    card = doc.tables[2]
    before = card._tbl.xml
    pasted = doc.add_table(rows=2, cols=2)
    pasted.cell(0, 0).text = "h"

    n = fixup.fix_content_tables(doc)

    assert n == 1
    assert card._tbl.xml == before
    assert pasted._tbl.tblPr.find(qn("w:jc")).get(qn("w:val")) == "center"


def test_masthead_grid_uses_layout_constants(tmp_path):
    from tme_template.layout import MASTHEAD_LEFT_WIDTH, MASTHEAD_RIGHT_WIDTH
    doc = Document(_starter(tmp_path))
    fixup.fix_masthead_grid(doc)
    cols = [int(c.get(qn("w:w"))) for c in doc.tables[0]._tbl.tblGrid.findall(qn("w:gridCol"))]
    assert cols == [MASTHEAD_LEFT_WIDTH.twips, MASTHEAD_RIGHT_WIDTH.twips]
