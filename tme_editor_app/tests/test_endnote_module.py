"""The EndNote resolver lives in the app now; a plain document passes through."""
from docx import Document


def test_resolver_is_importable_and_copies_a_plain_document(tmp_path):
    from endnote import resolve_endnote_citations

    src = tmp_path / "in.docx"
    dst = tmp_path / "out.docx"
    doc = Document()
    doc.add_paragraph("No EndNote fields here.")
    doc.save(src)

    resolve_endnote_citations(str(src), str(dst))

    assert dst.exists()
    assert [p.text for p in Document(dst).paragraphs] == ["No EndNote fields here."]
