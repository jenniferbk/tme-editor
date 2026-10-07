"""Figure and table captions must end up centered after fixup, even when the
pasted paragraph carries a direct alignment override or the starter's caption
styles predate the LEFT→CENTER switch."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import fixup
from tme_template.styles import (
    register_body_style, register_title_style,
    register_heading_styles, register_remaining_styles,
)


def _make_doc_with_styles():
    doc = Document()
    register_body_style(doc)
    register_title_style(doc)
    register_heading_styles(doc)
    register_remaining_styles(doc)
    return doc


def _direct_jc(p):
    pPr = p._p.find(qn("w:pPr"))
    return None if pPr is None else pPr.find(qn("w:jc"))


def test_strip_direct_formatting_removes_caption_alignment_override():
    doc = _make_doc_with_styles()
    fig = doc.add_paragraph("Figure 1. Left-aligned by the author.", style="TME Figure Caption")
    fig.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tab = doc.add_paragraph("Table 1. Also left-aligned.", style="TME Table Caption")
    tab.alignment = WD_ALIGN_PARAGRAPH.LEFT
    assert _direct_jc(fig) is not None and _direct_jc(tab) is not None

    fixup.strip_direct_formatting(doc)

    assert _direct_jc(fig) is None
    assert _direct_jc(tab) is None
    # With the override gone, the style's CENTER wins.
    assert fig.style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_strip_direct_formatting_keeps_body_alignment():
    """Image-holding body paragraphs are centered directly; leave them alone."""
    doc = _make_doc_with_styles()
    body = doc.add_paragraph("", style="TME Body")
    body.alignment = WD_ALIGN_PARAGRAPH.CENTER

    fixup.strip_direct_formatting(doc)

    assert body.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_update_styles_centers_caption_styles_from_older_starter():
    doc = _make_doc_with_styles()
    fc = doc.styles["TME Figure Caption"]
    tc = doc.styles["TME Table Caption"]
    fc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tc.paragraph_format.alignment = None

    fixup.update_styles(doc)

    assert fc.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert tc.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
