"""Page setup and cover footer: odd/even on, distances explicit, footer on both parities."""
from docx import Document
from docx.shared import Inches

from tme_template.cover_footer import add_cover_footer
from tme_template.page_setup import configure_page_setup


def _footer_text(footer):
    return " ".join(p.text for t in footer.tables for c in t._cells for p in c.paragraphs)


def test_odd_even_setting_is_on_and_distances_are_explicit():
    doc = Document()
    configure_page_setup(doc)
    assert doc.settings.odd_and_even_pages_header_footer is True
    s = doc.sections[0]
    assert s.header_distance == Inches(0.5)
    assert s.footer_distance == Inches(0.5)


def test_cover_footer_is_defined_for_even_pages_too():
    doc = Document()
    configure_page_setup(doc)
    sec = doc.sections[0]

    add_cover_footer(sec, citation="Cite me.")

    assert "HOW TO CITE" in _footer_text(sec.footer)
    assert "HOW TO CITE" in _footer_text(sec.even_page_footer)
    assert sec.even_page_footer.is_linked_to_previous is False
    xml = sec._sectPr.xml
    assert 'w:type="default"' in xml and 'w:type="even"' in xml
