"""Table widths: tcW (Word) and gridCol (LibreOffice) must agree; no blank DOI line."""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from tme_template.cover_page import AuthorEntry, CoverData, add_research_article_cover
from tme_template.layout import FULL_BLEED_WIDTH, MASTHEAD_LEFT_WIDTH, MASTHEAD_RIGHT_WIDTH
from tme_template.masthead import MastheadData, add_masthead
from tme_template.styles import register_title_style
from tme_template.tagline import add_tagline_strip

LOGO = str(Path(__file__).resolve().parents[2] / "assets" / "tme-logo.jpg")


def _grid(table):
    return [int(c.get(qn("w:w"))) for c in table._tbl.tblGrid.findall(qn("w:gridCol"))]


def _tcw(cell):
    return int(cell._tc.tcPr.find(qn("w:tcW")).get(qn("w:w")))


def _masthead(doi="doi.org/10.1/x"):
    return MastheadData(article_type="RESEARCH ARTICLE", volume=34, number=1, year=2026,
                        pages="1–24", doi=doi, issn_print="1062-9017",
                        issn_online="2331-4451", logo_path=LOGO)


def test_masthead_cells_match_grid_and_layout_constants():
    doc = Document()
    add_masthead(doc, _masthead())
    t = doc.tables[0]
    assert _grid(t) == [MASTHEAD_LEFT_WIDTH.twips, MASTHEAD_RIGHT_WIDTH.twips]
    assert [_tcw(c) for c in t.rows[0].cells] == _grid(t)


def test_masthead_without_doi_has_no_blank_line():
    doc = Document()
    add_masthead(doc, _masthead(doi=None))
    right = doc.tables[0].cell(0, 1)
    assert all(p.text.strip() for p in right.paragraphs)


def test_tagline_cell_is_full_bleed():
    doc = Document()
    add_tagline_strip(doc)
    t = doc.tables[0]
    assert _grid(t) == [FULL_BLEED_WIDTH.twips]
    assert _tcw(t.cell(0, 0)) == FULL_BLEED_WIDTH.twips


def test_author_card_widths_split_evenly_without_table_grid_style():
    doc = Document()
    register_title_style(doc)
    authors = [AuthorEntry(name=f"A{i}", affiliation_num=1, role=None, bio="b", headshot_path=None)
               for i in range(4)]
    add_research_article_cover(doc, CoverData(
        title="T", authors=authors, affiliations=["X"], dates={"Received": "d"},
        abstract="abs", keywords=["k"]))
    card = doc.tables[0]
    expected = Inches(7.5 / 4).twips
    assert _grid(card) == [expected] * 4
    assert [_tcw(c) for c in card.rows[0].cells] == [expected] * 4
    assert card.style is None or card.style.name != "Table Grid"
