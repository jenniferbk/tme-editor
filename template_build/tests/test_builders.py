"""Builders fail loudly on missing images and separate pages explicitly."""
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

from tme_template.front_matter import IssueInfo, add_formal_title_page, add_issue_cover_page
from tme_template.masthead import MastheadData, add_masthead

LOGO = str(Path(__file__).resolve().parents[2] / "assets" / "tme-logo.jpg")
LOGO_P = str(Path(__file__).resolve().parents[2] / "assets" / "tme-logo-portrait.jpg")


def test_masthead_with_missing_logo_raises():
    doc = Document()
    with pytest.raises(FileNotFoundError):
        add_masthead(doc, MastheadData(article_type="X", volume=1, number=1, year=2026, pages=None,
                                       doi=None, issn_print="1", issn_online="2",
                                       logo_path="/nonexistent/logo.jpg"))


def test_issue_cover_with_missing_portrait_raises():
    doc = Document()
    with pytest.raises(FileNotFoundError):
        add_issue_cover_page(doc, IssueInfo(volume=1, number=1, year=2026, season="Spring",
                                            cover_artist=None, portrait_logo_path="/nonexistent.jpg"))


def test_title_page_spacer_is_an_exact_line_height_not_space_before():
    doc = Document()
    add_formal_title_page(doc, IssueInfo(volume=1, number=1, year=2026, season="Spring",
                                         cover_artist=None, portrait_logo_path=LOGO_P))
    spacer = doc.paragraphs[0]
    assert spacer.paragraph_format.line_spacing == Pt(120)
    assert spacer.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert spacer.paragraph_format.space_before is None
