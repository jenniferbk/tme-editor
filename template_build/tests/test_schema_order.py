"""OOXML helpers must insert children in schema order."""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from schema_order import order_violations
from tme_template.oxml_helpers import (
    apply_bottom_rule, apply_top_rule, force_table_full_width,
    remove_cell_borders, set_cell_margins, set_cell_shading,
)


def test_paragraph_border_lands_before_spacing_indent_and_jc():
    doc = Document()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Inches(0.1)
    p.paragraph_format.left_indent = Inches(1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_bottom_rule(p, hex_color="BA0C2F", width_pt=1)
    apply_top_rule(p, hex_color="BA0C2F", width_pt=1)
    assert order_violations(doc.element.body) == []


def test_table_and_cell_properties_in_schema_order_whatever_the_call_order():
    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    c = t.cell(0, 0)
    set_cell_margins(c, top=1, bottom=1, left=1, right=1)   # tcMar first on purpose
    set_cell_shading(c, "FFFFFF")                           # shd must land before it
    remove_cell_borders(c)                                  # tcBorders before shd
    c.width = Inches(2)                                     # tcW first of all
    force_table_full_width(t, total_width_inches=8.5)
    assert order_violations(doc.element.body) == []
    tblPr = t._tbl.tblPr
    assert tblPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblLayout") is not None
